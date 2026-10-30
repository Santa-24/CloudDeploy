import docx
from collections import Counter

doc_path = r'C:/Users/santa/.gemini/antigravity/brain/4c4aa63f-7446-4c1c-be7b-91a46f45cc30/.user_uploaded/media_1793293080086_7523e202.docx'
doc = docx.Document(doc_path)

print('=== Sections & Margins ===')
for i, s in enumerate(doc.sections):
    print(f'Sec {i}: top={s.top_margin.pt}pt, bottom={s.bottom_margin.pt}pt, left={s.left_margin.pt}pt, right={s.right_margin.pt}pt')
    print(f'  header={s.header.paragraphs[0].text if s.header.paragraphs else None}')
    print(f'  footer={s.footer.paragraphs[0].text if s.footer.paragraphs else None}')

fonts = Counter()
sizes = Counter()
colors = Counter()
line_spacings = Counter()

for p in doc.paragraphs:
    pf = p.paragraph_format
    if pf.line_spacing:
        line_spacings[pf.line_spacing] += 1
    for r in p.runs:
        if r.font.name:
            fonts[r.font.name] += 1
        if r.font.size:
            sizes[r.font.size.pt] += 1
        if r.font.color and r.font.color.rgb:
            colors[str(r.font.color.rgb)] += 1

print('Fonts:', fonts.most_common(10))
print('Sizes:', sizes.most_common(10))
print('Colors:', colors.most_common(10))
print('Line spacings:', line_spacings.most_common(5))

print('\n=== Paragraph Styles ===')
style_counts = Counter(p.style.name for p in doc.paragraphs)
print(style_counts)

print('\n=== Sample Headings ===')
for p in doc.paragraphs:
    if p.style.name.startswith('Heading') or p.style.name in ['Title', 'Subtitle']:
        print(f'{p.style.name} [{p.runs[0].font.name if p.runs and p.runs[0].font.name else "def"} {p.runs[0].font.size.pt if p.runs and p.runs[0].font.size else "def"}]: {p.text[:60]}')

